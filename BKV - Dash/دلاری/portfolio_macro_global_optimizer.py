#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
=============================================================================
BKV Global Macro Multi-Asset Portfolio Optimizer (پورتفوی ماکرو دلاری جهانی)
=============================================================================
این اسکریپت عملکرد ۱۱ دارایی کلان بین‌المللی را به دلار آمریکا (USD)
در بزرگترین بازه مشترک تاریخی (۴۵ دوره ۳ ماهه / بیش از ۱۱ سال کامل از اواسط ۲۰۱۵)
با استفاده از الگوریتم‌های بهینه‌سازی ترکیبی (الگوریتم ژنتیک روی سیمپلکس،
نمونه‌برداری چندوجهی دبریشله و جستجوی محلی گام ۱٪) تحلیل، بهینه‌سازی و بک‌تست می‌کند.

دارایی‌های ۱۱گانه پورتفو:
۱. سکه طلا (Gold Coin) [GC=F]
۲. طلای ۲۴ عیار (Gold 24k) [GC=F]
۳. نقره جهانی (Silver) [SI=F]
۴. مس جهانی (Copper) [HG=F]
۵. روی جهانی (Zinc) [ZNC=F]
۶. سنگ آهن جهانی (Iron Ore) [TIO=F]
۷. نفت خام برنت / قیر (Brent) [BZ=F]
۸. بیت‌کوین (Bitcoin - BTC)
۹. اتریوم (Ethereum - ETH)
۱۰. شاخص سهام آمریکا (S&P 500 - SPX)
۱۱. تتر (Tether - USDT)

ویژگی‌های گزارش اکسل (۷ شیت کامل):
- شیت ۱: چرخش فصلی پویا (Dynamic Quarterly Rotation)
- شیت ۲: پورتفوی با وزن ثابت بیشترین بازدهی کل (Highest Return Fixed Portfolio)
- شیت ۳: پورتفوی بهینه ترکیبی بیشترین بازدهی با کمترین افت (Low Drawdown Optimal Portfolio)
- شیت ۴: رده‌بندی ۲۰۰ پورتفوی برتر بر اساس بازدهی (رنکینگ اصلی بازدهی) به همراه ستون Max Drawdown
- شیت ۵: رده‌بندی ۲۰۰ پورتفوی برتر بر اساس بیشترین بازدهی و کمترین Max Drawdown (شاخص کالمر)
- شیت ۶: عملکرد تک‌تک نمادهای ۱۱گانه در کل بازه ۱۱ ساله
- شیت ۷: پورتفوی سفارشی شما (Custom Portfolio Backtest) حاصل از بک‌تست تاریخی درصدهای دلخواه کاربر
- تضمین قطعی اینکه در تمامی ردیف‌ها و شیت‌های پورتفو، مجموع وزن‌ها با فرمول =SUM دقیقاً ۱۰۰٪ باشد.
=============================================================================
"""

import os
import sys
import time
import json
import argparse
import itertools
from datetime import datetime
import jdatetime
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# تنظیم انکودینگ ترمینال ویندوز برای نمایش صحیح متون فارسی
if sys.platform.startswith('win'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

# --- تنظیمات ---
BKV_JSON_NAME = "bkv_data.json"
MACRO_JSON_NAME = "macro_assets_data.json"
EXCEL_OUTPUT_NAME = "portfolio_macro_global_report.xlsx"
EXCEL_ALL_COMBINATIONS_NAME = "portfolio_macro_all_combinations.xlsx"

# دارایی‌های ۱۱گانه پروژه
MACRO_COMMODITIES = [
    ('gold_coin', 'سکه طلا (Gold Coin)', 'GC=F', 'A16207'),
    ('gold_995', 'طلای ۲۴ عیار (Gold 24k)', 'GC=F', 'B45309'),
    ('silver', 'نقره جهانی (Silver)', 'SI=F', '475569'),
    ('copper', 'مس جهانی (Copper)', 'HG=F', 'C2410C'),
    ('zinc', 'روی جهانی (Zinc)', 'ZNC=F', '64748B'),
    ('iron', 'سنگ آهن جهانی (Iron Ore)', 'TIO=F', '991B1B'),
    ('bitumen', 'قیر / نفت جهانی (Brent)', 'BZ=F', '1E293B'),
    ('btc', 'بیت‌کوین (Bitcoin - BTC)', 'BTC-USD', 'F59E0B'),
    ('eth', 'اتریوم (Ethereum - ETH)', 'ETH-USD', '6366F1'),
    ('spx', 'اس اند پی ۵۰۰ (S&P 500)', '^GSPC', '059669'),
    ('usdt', 'تتر (Tether - USDT)', 'USDT-USD', '0D9488'),
]


def load_datasets():
    """بارگذاری داده‌های کامودیتی و داده‌های کریپتو و سهام."""
    paths_bkv = [
        BKV_JSON_NAME,
        os.path.join(os.path.dirname(__file__), BKV_JSON_NAME),
        os.path.join(os.path.dirname(os.path.abspath(__file__)), BKV_JSON_NAME),
        os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', BKV_JSON_NAME),
        os.path.join('..', BKV_JSON_NAME),
    ]
    bkv_path = next((p for p in paths_bkv if os.path.exists(p)), None)
    if not bkv_path:
        print(f"❌ خطا: فایل داده '{BKV_JSON_NAME}' یافت نشد.")
        sys.exit(1)

    with open(bkv_path, 'r', encoding='utf-8') as f:
        bkv_data = json.load(f)

    paths_macro = [
        MACRO_JSON_NAME,
        os.path.join(os.path.dirname(__file__), MACRO_JSON_NAME),
        os.path.join(os.path.dirname(os.path.abspath(__file__)), MACRO_JSON_NAME),
    ]
    macro_path = next((p for p in paths_macro if os.path.exists(p)), None)
    if not macro_path:
        print(f"❌ خطا: فایل داده '{MACRO_JSON_NAME}' یافت نشد.")
        sys.exit(1)

    with open(macro_path, 'r', encoding='utf-8') as f:
        macro_data = json.load(f)

    return bkv_data, macro_data


def normalize_to_int_percent(weights_raw):
    """
    نرمال‌سازی وزن‌های اعشاری یا غیر ۱۰۰٪ به اعداد صحیح درصدی با مجموع دقیق ۱۰۰٪
    با استفاده از روش بزرگ‌ترین باقیمانده (Largest Remainder / Hamilton Method).
    """
    w_float = np.array(weights_raw, dtype=float)
    s = np.sum(w_float)
    if s <= 0:
        res = np.zeros(len(w_float), dtype=int)
        res[0] = 100
        return res
    scaled = (w_float / s) * 100.0
    floor_parts = np.floor(scaled).astype(int)
    diff = 100 - np.sum(floor_parts)
    remainders = scaled - floor_parts
    order = np.argsort(remainders)[::-1]
    for i in range(diff):
        floor_parts[order[i % len(floor_parts)]] += 1
    return floor_parts


def parse_cli_weights(weights_str, cid_list):
    """پردازش رشته ورودی خط فرمان برای درصدهای سفارشی."""
    parts = [p.strip() for p in weights_str.split(',') if p.strip()]
    weights = np.zeros(len(cid_list), dtype=float)
    if any('=' in p for p in parts):
        mapping = {}
        for p in parts:
            if '=' in p:
                k, v = p.split('=', 1)
                mapping[k.strip().lower()] = float(v.replace('%', '').strip())
        for i, cid in enumerate(cid_list):
            for k, val in mapping.items():
                if k in [cid.lower(), str(i), str(i + 1)]:
                    weights[i] = val
    else:
        for i, p in enumerate(parts[:len(cid_list)]):
            weights[i] = float(p.replace('%', '').strip())
    return normalize_to_int_percent(weights)


def prompt_custom_weights(cid_list, cname_list):
    """دریافت تعاملی درصد وزن هر یک از ۱۱ دارایی از کاربر در محیط ترمینال."""
    print("\n" + "═" * 80)
    print("🎯 دریافت ترکیب درصد پورتفوی سفارشی شما برای بک‌تست تاریخی:")
    print("═" * 80)
    print("💡 برای هر دارایی درصد مورد نظر (بین ۰ تا ۱۰۰) را وارد کنید.")
    print("💡 در صورت تمایل به وزن ۰٪، کافیست کلید اینتر (Enter) را بزنید.\n")

    weights = []
    current_sum = 0
    for idx, (cid, cname) in enumerate(zip(cid_list, cname_list), 1):
        rem = max(0, 100 - current_sum)
        while True:
            prompt_str = f"  [{idx:02d}/11] {cname:<26} (حداکثر باقی‌مانده: {rem}%): "
            try:
                raw_in = input(prompt_str).strip()
            except (EOFError, KeyboardInterrupt):
                print("\n\n⚠️ ورود اطلاعات متوقف شد. از مقادیر پیش‌فرض استفاده می‌شود.")
                return None
            if raw_in == '':
                val = 0
                weights.append(val)
                break
            try:
                val = int(round(float(raw_in.replace('%', ''))))
                if val < 0:
                    print("     ❌ لطفاً عدد نامنفی وارد کنید.")
                    continue
                weights.append(val)
                current_sum += val
                break
            except ValueError:
                print("     ❌ مقدار وارد شده نامعتبر است. لطفاً عدد وارد کنید.")

    tot = sum(weights)
    if tot == 100:
        print(f"\n✅ مجموع وزن‌ها دقیقاً ۱۰۰٪ شد.")
        return np.array(weights, dtype=int)
    else:
        print(f"\n⚠️ مجموع درصدهای وارد شده: {tot}% (باید ۱۰۰٪ باشد).")
        try:
            fix = input("آیا مایلید به صورت خودکار نرمالیزه شود تا دقیقاً ۱۰۰٪ شود؟ [Y/n]: ").strip().lower()
        except Exception:
            fix = 'y'
        if fix in ['', 'y', 'yes', 'بله']:
            norm_w = normalize_to_int_percent(weights)
            print("✅ درصدها به شرح زیر نرمالیزه شدند:")
            for cname, w in zip(cname_list, norm_w):
                if w > 0:
                    print(f"   • {cname}: {w}%")
            return norm_w
        else:
            return None


def get_custom_portfolio_weights(args, cid_list, cname_list):
    """
    بررسی و تامین وزن‌های پورتفوی سفارشی بر اساس آرگومان‌ها، پرامپت تعاملی، یا پیش‌فرض.
    """
    # ۱. آرگومان مستقیم --weights
    if args.weights:
        w_int = parse_cli_weights(args.weights, cid_list)
        print(f"\n✅ درصدهای سفارشی از خط فرمان دریافت شد (مجموع: {sum(w_int)}%).")
        return w_int, True

    # ۲. حالت تعاملی (اگر پرچم --custom زده شده باشد یا ترمینال زنده باشد و پرچم no-interactive نباشد)
    is_interactive = args.custom or (sys.stdin.isatty() and not args.no_interactive)
    if is_interactive and args.custom:
        w_int = prompt_custom_weights(cid_list, cname_list)
        if w_int is not None:
            return w_int, True

    if is_interactive and not args.no_interactive:
        print("\n" + "═" * 80)
        print("🎯 بخش بک‌تست پورتفوی پیشنهادی / سفارشی (Custom Portfolio Backtester)")
        print("═" * 80)
        try:
            ans = input("آیا مایلید ترکیب درصد پیشنهادی خود را وارد کنید تا در طول کل تاریخچه ۱۱ ساله بک‌تست شود؟ [y/N]: ").strip().lower()
        except Exception:
            ans = 'n'
        if ans in ['y', 'yes', 'بله']:
            w_int = prompt_custom_weights(cid_list, cname_list)
            if w_int is not None:
                return w_int, True

    # ۳. پورتفوی پیش‌فرض الگو در صورت عدم ورود کاربر
    # پورتفوی متوازن نمونه: ۳۰٪ طلای ۲۴ عیار، ۵٪ نفت برنت، ۱۰٪ بیت‌کوین، ۱۵٪ اتریوم، ۲۰٪ S&P 500، ۲۰٪ تتر
    default_w = np.array([0, 30, 0, 0, 0, 0, 5, 10, 15, 20, 20], dtype=int)
    return default_w, False


def backtest_single_portfolio(w11_int, R, P_start, P_end, periods_data):
    """
    محاسبه کامل عملکرد یک چینش وزنی خاص در کل ۴۵ دوره فصلی.
    """
    T = len(periods_data)
    w11 = w11_int / 100.0
    records = []
    c_mult = 1.0
    cap_arr = []
    p_rets = []

    for t in range(T):
        p_ret = float(R[t] @ w11)
        c_mult *= (1.0 + p_ret)
        cap_arr.append(c_mult)
        p_rets.append(p_ret)

        records.append({
            'period': periods_data[t],
            'weights': w11,
            'period_ret': p_ret,
            'cum_capital': 1000.0 * c_mult,
            'cum_return': c_mult - 1.0,
        })

    cap_arr = np.array(cap_arr)
    p_rets = np.array(p_rets)
    peaks = np.maximum.accumulate(cap_arr)
    dds = (cap_arr - peaks) / peaks
    mdd = float(np.min(dds))
    cum_ret = float(c_mult - 1.0)
    years = T / 4.0
    cagr = float((1.0 + max(cum_ret, -0.9999)) ** (1.0 / years) - 1.0)
    calm = float(cagr / max(abs(mdd), 0.001)) if abs(mdd) > 0.001 else 0.0
    mean_q = float(np.mean(p_rets))
    std_q = float(np.std(p_rets))
    sh = float(mean_q / (std_q + 1e-6))
    worst_q = float(np.min(p_rets))
    best_q = float(np.max(p_rets))
    win_rate = float(np.mean(p_rets > 0))

    summary = {
        'weights': w11,
        'weights_int': w11_int,
        'cum_return': cum_ret,
        'cagr': cagr,
        'max_drawdown': mdd,
        'mean_return': mean_q,
        'std_return': std_q,
        'sharpe': sh,
        'calmar': calm,
        'worst_quarter': worst_q,
        'best_quarter': best_q,
        'win_rate': win_rate,
        'final_capital_1k': 1000.0 * c_mult,
    }
    return records, summary


def fast_vectorized_eval(W_int, R_dist, T):
    """ارزیابی برداری با سرعت بالا برای ماتریس وزن‌های کاندیدا."""
    W = W_int / 100.0
    P = W @ R_dist.T  # شکل: (K, T)
    cap = np.cumprod(1.0 + P, axis=1)
    peaks = np.maximum(np.maximum.accumulate(cap, axis=1), 1e-9)
    dds = (cap - peaks) / peaks
    mdds = np.clip(np.min(dds, axis=1), -1.0, 0.0)
    cums = cap[:, -1] - 1.0
    cagr = np.where(cums > -0.9999, (1.0 + np.maximum(cums, -0.9999)) ** (1.0 / (T / 4.0)) - 1.0, -1.0)
    safe_mdd = np.maximum(np.abs(mdds), 1e-4)
    calm = np.where(np.abs(mdds) > 0.001, cagr / safe_mdd, 0.0)
    calm = np.nan_to_num(calm, nan=0.0, posinf=0.0, neginf=0.0)
    means = np.mean(P, axis=1)
    stds = np.std(P, axis=1)
    sharpes = means / (stds + 1e-6)
    worsts = np.min(P, axis=1)
    bests = np.max(P, axis=1)
    return cums, cagr, mdds, calm, means, worsts, bests, sharpes


def optimize_multi_asset_portfolios(R, distinct_indices, T):
    """
    موتور بهینه‌سازی ترکیبی چنددارایی (Hybrid Multi-Asset Optimizer):
    ۱. گرید سیستماتیک ۱ دارایی (۱۰ حالت)
    ۲. گرید سیستماتیک ۲ دارایی با گام ۱٪ (۴,۴۵۵ حالت)
    ۳. گرید سیستماتیک ۳ دارایی با گام ۲٪ (۱۴۱,۱۲۰ حالت)
    ۴. الگوریتم ژنتیک روی سیمپلکس (Simplex GA) روی ۴ تا ۱۰ دارایی:
       - جمعیت: ۱۰,۰۰۰ فرد در ۳۵ نسل
       - اپراتورهای جهش تبادل وزن (Pairwise Mutation) با گام ۱٪، ۲٪، ۳٪، ۵٪
       - اپراتور تقاطع محدب تصادفی روی سیمپلکس
       - تزریق نمونه‌برداری دبریشله چند مقیاسی (Multi-Alpha Dirichlet Immigration)
    ۵. جستجوی صعودی محلی گام ۱٪ (1% Local Hill Climbing Polish) روی ۲۰۰ کاندیدای برتر
    ۶. گردآوری تمام پورتفوهای یکتا در یک استخر عظیم و استخراج ۲۰۰ تای برتر
    """
    R_dist = R[:, distinct_indices]
    N_dist = len(distinct_indices)

    print("\n" + "=" * 85)
    print("🧬 آغاز فرآیند بهینه‌سازی ترکیبی (الگوریتم ژنتیک + دبریشله + جستجوی محلی گام ۱٪)")
    print("=" * 85)

    w_pool = set()

    # ۱. گرید سیستماتیک ۱ دارایی (۱۰۰٪)
    for i in range(N_dist):
        w = [0] * N_dist
        w[i] = 100
        w_pool.add(tuple(w))

    # ۲. گرید سیستماتیک ۲ دارایی با گام ۱٪
    for i in range(N_dist):
        for j in range(i + 1, N_dist):
            for p in range(1, 100):
                w = [0] * N_dist
                w[i] = p
                w[j] = 100 - p
                w_pool.add(tuple(w))

    # ۳. گرید سیستماتیک ۳ دارایی با گام ۲٪
    for i in range(N_dist):
        for j in range(i + 1, N_dist):
            for k in range(j + 1, N_dist):
                for p1 in range(2, 98, 2):
                    for p2 in range(2, 100 - p1, 2):
                        w = [0] * N_dist
                        w[i] = p1
                        w[j] = p2
                        w[k] = 100 - p1 - p2
                        w_pool.add(tuple(w))

    print(f"  🔹 فضای گرید مقدماتی (۱، ۲ و ۳ دارایی): {len(w_pool):,} پورتفوی یکتا تولید شد.")

    # ۴. الگوریتم ژنتیک چنددارایی روی سیمپلکس (Simplex Genetic Algorithm)
    pop_size = 10000
    generations = 35
    alphas = [0.05, 0.1, 0.2, 0.5, 1.0, 2.0]

    # تولید جمعیت اولیه دبریشله چندآلفا
    init_pop = []
    for _ in range(pop_size):
        a = np.random.choice(alphas)
        d = np.random.dirichlet([a] * N_dist)
        di = np.round(d * 100).astype(int)
        di[np.argmax(di)] += (100 - np.sum(di))
        init_pop.append(di)
    pop = np.array(init_pop)

    print(f"  🧬 اجرای الگوریتم ژنتیک روی ترکیب‌های ۴ تا ۱۰ دارایی ({generations} نسل با جمعیت {pop_size:,})...")

    for gen in range(generations):
        cums, cagr, mdds, calm, _, _, _, _ = fast_vectorized_eval(pop, R_dist, T)
        for row in pop:
            w_pool.add(tuple(row.tolist()))

        # رتبه‌بندی نخبگان: ۱۵۰۰ نخبه برتر کالمر + ۵۰۰ نخبه برتر بازدهی کل
        ord_calm = np.argsort(calm)[::-1]
        ord_ret = np.argsort(cums)[::-1]
        elites_calm = pop[ord_calm[:1500]]
        elites_ret = pop[ord_ret[:500]]
        elites = np.vstack([elites_calm, elites_ret])

        # جهش دوتایی (Pairwise Mutation)
        mutants = elites.copy()
        for r in range(len(mutants)):
            pos = np.where(mutants[r] > 0)[0]
            if len(pos) > 0:
                i_f = np.random.choice(pos)
                i_t = np.random.choice(N_dist)
                st = min(int(np.random.choice([1, 2, 3, 5])), mutants[r, i_f])
                mutants[r, i_f] -= st
                mutants[r, i_t] += st

        # تقاطع محدب (Convex Simplex Crossover)
        p1 = elites[np.random.choice(len(elites), 4000)]
        p2 = elites[np.random.choice(len(elites), 4000)]
        al = np.random.uniform(0.1, 0.9, (4000, 1))
        cross = np.round(al * p1 + (1 - al) * p2).astype(int)
        diff = 100 - np.sum(cross, axis=1)
        for r in range(4000):
            cross[r, np.argmax(cross[r])] += diff[r]

        # تزریق مهاجران جدید دبریشله (Dirichlet Immigration) برای تنوع ژنتیکی
        imm_list = []
        for _ in range(2000):
            a = np.random.choice(alphas)
            d = np.random.dirichlet([a] * N_dist)
            di = np.round(d * 100).astype(int)
            di[np.argmax(di)] += (100 - np.sum(di))
            imm_list.append(di)
        imm = np.array(imm_list)

        pop = np.vstack([elites, mutants, cross, imm])

    print(f"  🔍 استخر نهایی پورتفوها: {len(w_pool):,} ترکیب یکتا کشف و گردآوری شد.")

    # ۵. جستجوی صعودی محلی گام ۱٪ (Local 1% Hill Climbing Polish)
    print("  🏔️ اجرای جستجوی صعودی محلی گام ۱٪ (Local Hill Climbing) برای قله‌های بهینه...")
    all_W_pre = np.array(list(w_pool), dtype=int)
    cums_pre, cagr_pre, mdds_pre, calm_pre, _, _, _, _ = fast_vectorized_eval(all_W_pre, R_dist, T)

    top_calm_indices = np.argsort(calm_pre)[::-1][:150]
    top_ret_indices = np.argsort(cums_pre)[::-1][:50]
    polish_seeds = np.vstack([all_W_pre[top_calm_indices], all_W_pre[top_ret_indices]])

    for seed in polish_seeds:
        for i in range(N_dist):
            if seed[i] > 0:
                for j in range(N_dist):
                    if i != j:
                        w_cand = seed.copy()
                        w_cand[i] -= 1
                        w_cand[j] += 1
                        w_pool.add(tuple(w_cand.tolist()))

    print(f"  ✨ مجموع کل پورتفوهای ارزیابی‌شده در گرید و الگوریتم: {len(w_pool):,} ترکیب مختلف.")

    # ۶. ارزیابی نهایی کل استخر پورتفوها
    all_W = np.array(list(w_pool), dtype=int)
    cums, cagr, mdds, calm, means, worsts, bests, sharpes = fast_vectorized_eval(all_W, R_dist, T)

    # تابع نگاشت وزن ۱۰تایی به ۱۱تایی
    def map_to_11(w10):
        w11 = np.zeros(len(MACRO_COMMODITIES), dtype=float)
        w11[0] = 0.0  # سکه
        for idx_d, orig_idx in enumerate(distinct_indices):
            w11[orig_idx] = w10[idx_d] / 100.0
        return w11

    # استخراج ۲۰۰ پورتفوی برتر بر اساس بازدهی کل (رنکینگ اصلی)
    order_by_ret = np.argsort(cums)[::-1]
    top_200_by_return = []
    for idx in order_by_ret[:200]:
        top_200_by_return.append({
            'weights': map_to_11(all_W[idx]),
            'weights_int': all_W[idx],
            'cum_return': cums[idx],
            'cagr': cagr[idx],
            'mean_return': means[idx],
            'max_drawdown': mdds[idx],
            'worst_quarter': worsts[idx],
            'best_quarter': bests[idx],
            'sharpe': sharpes[idx],
            'calmar': calm[idx],
        })

    # استخراج ۲۰۰ پورتفوی برتر بر اساس شاخص کالمر (بازدهی بالا و کمترین افت)
    order_by_calmar = np.argsort(calm)[::-1]
    top_200_by_calmar = []
    for idx in order_by_calmar[:200]:
        top_200_by_calmar.append({
            'weights': map_to_11(all_W[idx]),
            'weights_int': all_W[idx],
            'cum_return': cums[idx],
            'cagr': cagr[idx],
            'mean_return': means[idx],
            'max_drawdown': mdds[idx],
            'worst_quarter': worsts[idx],
            'best_quarter': bests[idx],
            'sharpe': sharpes[idx],
            'calmar': calm[idx],
        })

    print("=" * 85)
    print("✅ بهینه‌سازی ترکیبی با موفقیت تکمیل شد و ۲۰۰ پورتفوی برتر هر دو رده‌بندی استخراج گردیدند.")
    print("=" * 85)

    full_db_data = (all_W, cums, cagr, mdds, calm, worsts, bests, sharpes, distinct_indices)
    return top_200_by_return, top_200_by_calmar, full_db_data


def run_macro_global_optimization(args=None):
    if args is None:
        args = argparse.Namespace(custom=False, weights=None, no_interactive=False)

    bkv_data, macro_data = load_datasets()

    comm_ids = [c[0] for c in MACRO_COMMODITIES]
    comm_names = [c[1] for c in MACRO_COMMODITIES]
    N = len(comm_ids)

    dates = macro_data['dates']
    T = len(dates) - 1

    print("\n" + "=" * 85)
    print("🌍 ابرپورتفوی ماکرو دلاری بازارهای جهانی (۱۱ دارایی - ۴۵ دوره ۳ ماهه)")
    print("=" * 85)
    print(f"📅 بازه زمانی مشترک: {dates[0]} تا {dates[-1]} ({T} دوره ۳ ماهه کامل / ۱۱ سال و ۳ ماه)")
    print(f"💎 دارایی‌های فعال: طلا، سکه، نقره، مس، روی، سنگ آهن، نفت برنت، بیت‌کوین، اتریوم، S&P 500 و تتر")
    print("=" * 85)

    # استخراج دوره‌ها و ماتریس‌های قیمت و بازدهی R
    periods_data = []
    R = np.zeros((T, N))
    P_start = np.zeros((T, N))
    P_end = np.zeros((T, N))

    for t in range(T):
        d_s = dates[t]
        d_e = dates[t + 1]

        dt_s = pd.to_datetime(d_s)
        dt_e = pd.to_datetime(d_e)

        jd_s = jdatetime.date.fromgregorian(date=dt_s.date())
        jd_e = jdatetime.date.fromgregorian(date=dt_e.date())

        cal_days = (dt_e - dt_s).days

        periods_data.append({
            'period_num': t + 1,
            'greg_start': d_s,
            'greg_end': d_e,
            'jalali_start': jd_s.strftime('%Y/%m/%d'),
            'jalali_end': jd_e.strftime('%Y/%m/%d'),
            'days': cal_days,
        })

        # ۷ نماد اول از bkv_data
        for i, (cid, _, _, _) in enumerate(MACRO_COMMODITIES[:7]):
            pts = {p['date'][:7]: float(p['global_usd']) for p in bkv_data[cid]['historical']}
            p1 = pts[d_s[:7]]
            p2 = pts[d_e[:7]]
            P_start[t, i] = p1
            P_end[t, i] = p2
            R[t, i] = (p2 - p1) / p1

        # ۴ نماد ماکرو از macro_data: btc, eth, spx, usdt
        for idx_m, mk in enumerate(['btc', 'eth', 'spx', 'usdt'], 7):
            p1 = float(macro_data[mk][t])
            p2 = float(macro_data[mk][t + 1])
            P_start[t, idx_m] = p1
            P_end[t, idx_m] = p2
            R[t, idx_m] = (p2 - p1) / p1 if p1 > 0 else 0.0

    # ۱. عملکرد تک‌تک نمادهای منفرد (۱۰۰٪ در هر نماد)
    single_asset_stats = []
    print("\n📊 عملکرد تک‌تک نمادها در بازار جهانی به دلار (۴۵ دوره ۳ ماهه):")
    print("-" * 92)
    print(f"{'نام دارایی':<27} {'تیکر':<10} {'قیمت اول':<12} {'قیمت آخر':<12} {'سود کل (٪)':<14} {'Max DD':<10} {'شارپ'}")
    print("-" * 92)

    for i, (cid, cname, ticker, _) in enumerate(MACRO_COMMODITIES):
        cum = np.prod(1.0 + R[:, i]) - 1.0
        m_q = np.mean(R[:, i])
        s_q = np.std(R[:, i])
        sh = m_q / (s_q + 1e-6)
        worst_q = np.min(R[:, i])
        best_q = np.max(R[:, i])
        win = np.mean(R[:, i] > 0)
        p1 = P_start[0, i]
        p2 = P_end[-1, i]

        cap_i = np.cumprod(1.0 + R[:, i])
        peak_i = np.maximum.accumulate(cap_i)
        dd_i = (cap_i - peak_i) / peak_i
        mdd_i = np.min(dd_i)

        stat = {
            'cid': cid,
            'name': cname,
            'ticker': ticker,
            'p_start': p1,
            'p_end': p2,
            'cum_return': cum,
            'mean_q': m_q,
            'std_q': s_q,
            'sharpe': sh,
            'worst_q': worst_q,
            'best_q': best_q,
            'max_drawdown': mdd_i,
            'win_rate': win,
        }
        single_asset_stats.append(stat)
        print(f"{cname:<27} {ticker:<10} ${p1:<11.2f} ${p2:<11.2f} {cum*100:>+11.2f}% {mdd_i*100:>8.2f}% {sh:>7.2f}")
    print("-" * 92)

    # ۲. مدل اول: چرخش فصلی پویا (Dynamic Quarterly Rotation)
    dyn_records = []
    cum_mult_dyn = 1.0
    for t in range(T):
        best_i = int(np.argmax(R[t]))
        p_ret = float(R[t, best_i])
        cum_mult_dyn *= (1.0 + p_ret)

        w_t = np.zeros(N)
        w_t[best_i] = 1.0

        dyn_records.append({
            'period': periods_data[t],
            'weights': w_t,
            'chosen_idx': best_i,
            'chosen_name': comm_names[best_i],
            'period_ret': p_ret,
            'cum_capital': 1000.0 * cum_mult_dyn,
            'cum_return': cum_mult_dyn - 1.0,
        })

    dyn_period_rets = np.array([r['period_ret'] for r in dyn_records])
    dyn_cap = np.cumprod(1.0 + dyn_period_rets)
    dyn_peak = np.maximum.accumulate(dyn_cap)
    dyn_mdd = np.min((dyn_cap - dyn_peak) / dyn_peak)

    dyn_summary = {
        'cum_return': cum_mult_dyn - 1.0,
        'mean_return': np.mean(dyn_period_rets),
        'std_return': np.std(dyn_period_rets),
        'sharpe': np.mean(dyn_period_rets) / (np.std(dyn_period_rets) + 1e-6),
        'worst_period': np.min(dyn_period_rets),
        'best_period': np.max(dyn_period_rets),
        'max_drawdown': dyn_mdd,
        'win_rate': np.mean(dyn_period_rets > 0),
    }

    print("\n" + "=" * 85)
    print("🔄 مدل ۱: چرخش فصلی پویا (Dynamic Quarterly Rotation - ۱۱ دارایی)")
    print("=" * 85)
    print(f"🚀 مجموع بازدهی دلاری کل:    {dyn_summary['cum_return']*100:+.2f}%")
    print(f"📈 میانگین سود دلاری هر فصل: {dyn_summary['mean_return']*100:+.2f}%")
    print(f"📉 بیشترین افت (Max DD):     {dyn_summary['max_drawdown']*100:.2f}%")
    print(f"⚖️ شاخص شارپ دلاری:          {dyn_summary['sharpe']:.2f}")
    print(f"🟢 بهترین فصل دلاری:         {dyn_summary['best_period']*100:+.2f}%")
    print(f"🔴 بدترین فصل دلاری:         {dyn_summary['worst_period']*100:+.2f}%")
    print("=" * 85)

    # ۳. اجرای بهینه‌سازی الگوریتم ژنتیک و دبریشله برای استخراج ۲۰۰ پورتفوی برتر
    distinct_indices = [1, 2, 3, 4, 5, 6, 7, 8, 9, 10]
    top_200_by_return, top_200_by_calmar, full_db_data = optimize_multi_asset_portfolios(R, distinct_indices, T)

    best_ret_portfolio = top_200_by_return[0]
    best_balanced_portfolio = top_200_by_calmar[0]

    # محاسبه مسیر فصلی پورتفوی با بیشترین بازدهی
    fixed_ret_records = []
    c_mult = 1.0
    for t in range(T):
        p_ret = float(R[t] @ best_ret_portfolio['weights'])
        c_mult *= (1.0 + p_ret)
        fixed_ret_records.append({
            'period': periods_data[t],
            'weights': best_ret_portfolio['weights'],
            'period_ret': p_ret,
            'cum_capital': 1000.0 * c_mult,
            'cum_return': c_mult - 1.0,
        })

    # محاسبه مسیر فصلی پورتفوی بهینه بازدهی-افت (Balanced Low Drawdown)
    fixed_bal_records = []
    c_mult_b = 1.0
    for t in range(T):
        p_ret = float(R[t] @ best_balanced_portfolio['weights'])
        c_mult_b *= (1.0 + p_ret)
        fixed_bal_records.append({
            'period': periods_data[t],
            'weights': best_balanced_portfolio['weights'],
            'period_ret': p_ret,
            'cum_capital': 1000.0 * c_mult_b,
            'cum_return': c_mult_b - 1.0,
        })

    print("\n" + "=" * 85)
    print("🏆 مدل ۲: پورتفوی با بیشترین بازدهی کل (Highest Return Fixed Portfolio)")
    print("=" * 85)
    for idx, w in enumerate(best_ret_portfolio['weights']):
        if w > 0:
            print(f"  ⭐ {comm_names[idx]}: {int(round(w * 100))}%")
    print(f"  📐 جمع وزن‌های پورتفو: {sum(best_ret_portfolio['weights']) * 100:.1f}%")
    print("-" * 85)
    print(f"🚀 مجموع بازدهی دلاری کل:    {best_ret_portfolio['cum_return']*100:+.2f}%")
    print(f"📈 بازده سالانه دلاری (CAGR): {best_ret_portfolio['cagr']*100:+.2f}%")
    print(f"📉 بیشترین افت (Max DD):     {best_ret_portfolio['max_drawdown']*100:.2f}%")
    print(f"🔴 بیشترین ریزش یک فصل:     {best_ret_portfolio['worst_quarter']*100:.2f}%")
    print(f"⚖️ شاخص شارپ دلاری:          {best_ret_portfolio['sharpe']:.2f}")
    print("=" * 85)

    print("\n" + "=" * 85)
    print("🛡️ مدل ۳: پورتفوی بهینه بازدهی بالا با افت بسیار کم (Low Drawdown Optimal - رتبه ۱ کالمر)")
    print("=" * 85)
    for idx, w in enumerate(best_balanced_portfolio['weights']):
        if w > 0:
            print(f"  ⭐ {comm_names[idx]}: {int(round(w * 100))}%")
    print(f"  📐 جمع وزن‌های پورتفو: {sum(best_balanced_portfolio['weights']) * 100:.1f}%")
    print("-" * 85)
    print(f"🚀 مجموع بازدهی دلاری کل:    {best_balanced_portfolio['cum_return']*100:+.2f}%")
    print(f"📈 بازده سالانه دلاری (CAGR): {best_balanced_portfolio['cagr']*100:+.2f}%")
    print(f"📉 بیشترین افت (Max DD):     {best_balanced_portfolio['max_drawdown']*100:.2f}%")
    print(f"🔴 بیشترین ریزش یک فصل:     {best_balanced_portfolio['worst_quarter']*100:.2f}%")
    print(f"🎯 شاخص کارایی کالمر:        {best_balanced_portfolio['calmar']:.2f}")
    print(f"⚖️ شاخص شارپ دلاری:          {best_balanced_portfolio['sharpe']:.2f}")
    print("=" * 85)

    # ۴. مدل چهارم: بک‌تست پورتفوی سفارشی کاربر (Custom Portfolio Backtest)
    custom_w_int, is_user_custom = get_custom_portfolio_weights(args, comm_ids, comm_names)
    custom_records, custom_summary = backtest_single_portfolio(custom_w_int, R, P_start, P_end, periods_data)

    print("\n" + "=" * 85)
    if is_user_custom:
        print("🎯 مدل ۴: بک‌تست تاریخی پورتفوی سفارشی پیشنهادی شما")
    else:
        print("🎯 مدل ۴: بک‌تست پورتفوی نمونه (برای ورود درصد شخصی، اسکریپت را با --custom اجرا کنید)")
    print("=" * 85)
    for idx, w_int in enumerate(custom_w_int):
        if w_int > 0:
            print(f"  🔹 {comm_names[idx]}: {w_int}%")
    print(f"  📐 جمع وزن‌های پورتفوی شما: {sum(custom_w_int)}%")
    print("-" * 85)
    print(f"🚀 سود کل دلاری در ۱۱ سال:    {custom_summary['cum_return']*100:+.2f}%")
    print(f"📈 نرخ بازده مرکب سالانه (CAGR): {custom_summary['cagr']*100:+.2f}%")
    print(f"📉 بیشترین افت تاریخی (Max DD): {custom_summary['max_drawdown']*100:.2f}%")
    print(f"🔴 عمیق‌ترین افت یک فصل:       {custom_summary['worst_quarter']*100:.2f}%")
    print(f"🟢 پربازده‌ترین فصل پورتفو:     {custom_summary['best_quarter']*100:+.2f}%")
    print(f"⚖️ شاخص شارپ دلاری:            {custom_summary['sharpe']:.2f}")
    print(f"🎯 شاخص کارایی کالمر:          {custom_summary['calmar']:.2f}")
    print(f"💰 رشد ۱,۰۰۰ دلار سرمایه اولیه: ${custom_summary['final_capital_1k']:,.2f}")
    print("=" * 85)

    # ۵. صدور فایل اکسل جامع (شامل ۷ شیت کامل)
    export_macro_excel(
        periods_data,
        R,
        P_start,
        P_end,
        dyn_records,
        dyn_summary,
        fixed_ret_records,
        best_ret_portfolio,
        fixed_bal_records,
        best_balanced_portfolio,
        top_200_by_return,
        top_200_by_calmar,
        single_asset_stats,
        custom_records,
        custom_summary,
        is_user_custom,
        EXCEL_OUTPUT_NAME
    )
    print(f"\n🎉 فایل اکسل جامع پورتفوی ماکرو دلاری با موفقیت ذخیره شد:\n👉 {os.path.abspath(EXCEL_OUTPUT_NAME)}")

    # ۶. صدور فایل اکسل تک‌شیته بانک جامع پورتفوها (تمام حالات ارزیابی‌شده)
    export_full_database_excel(full_db_data, EXCEL_ALL_COMBINATIONS_NAME)


def export_full_database_excel(full_db_data, filename):
    """
    صدور فایل اکسل تک‌شیته جامع شامل تمامی حالات پورتفوی ارزیابی‌شده در الگوریتم
    با استفاده از حالت Write-Only و استریمینگ سریع openpyxl.
    """
    all_W, cums, cagr, mdds, calm, worsts, bests, sharpes, distinct_indices = full_db_data
    total_count = len(all_W)
    print(f"\n📑 در حال ایجاد فایل اکسل تک‌شیته بانک جامع پورتفوها ({total_count:,} ردیف)...")
    t0 = time.time()

    order = np.argsort(calm)[::-1]

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(title="تمام حالات پورتفو")
    ws.views.sheetView[0].rightToLeft = True

    headers = [
        "رتبه (کالمر)",
        "سکه طلا (٪)", "طلا ۲۴ (٪)", "نقره (٪)", "مس (٪)", "روی (٪)", "سنگ آهن (٪)", "قیر/نفت (٪)",
        "بیت‌کوین (٪)", "اتریوم (٪)", "S&P 500 (٪)", "تتر (٪)", "جمع وزن‌ها (٪)",
        "شاخص کالمر (Calmar)", "بیشترین افت کل (Max DD ٪)", "بازده مرکب سالانه (CAGR ٪)",
        "سود کل دلاری (٪)", "بیشترین ریزش یک فصل (٪)", "پربازده‌ترین فصل (٪)", "شاخص شارپ"
    ]
    ws.append(headers)

    for rank_i, idx in enumerate(order, 1):
        w10 = all_W[idx]
        row = [
            rank_i,
            0, int(w10[0]), int(w10[1]), int(w10[2]), int(w10[3]), int(w10[4]),
            int(w10[5]), int(w10[6]), int(w10[7]), int(w10[8]), int(w10[9]),
            100,
            round(float(calm[idx]), 2),
            round(float(mdds[idx] * 100), 2),
            round(float(cagr[idx] * 100), 2),
            round(float(cums[idx] * 100), 2),
            round(float(worsts[idx] * 100), 2),
            round(float(bests[idx] * 100), 2),
            round(float(sharpes[idx]), 2),
        ]
        ws.append(row)

    try:
        wb.save(filename)
        print(f"🎉 فایل اکسل تک‌شیته جامع ({total_count:,} ردیف) با موفقیت ذخیره شد ({time.time()-t0:.1f} ثانیه):\n👉 {os.path.abspath(filename)}")
    except PermissionError:
        alt_f = filename.replace('.xlsx', '_v2.xlsx')
        wb.save(alt_f)
        print(f"⚠️ فایل قفل بود؛ در فایل جایگزین ذخیره شد:\n👉 {os.path.abspath(alt_f)}")



def export_macro_excel(periods_data, R, P_start, P_end, dyn_records, dyn_summary,
                       fixed_ret_records, best_ret, fixed_bal_records, best_bal,
                       top_200_ret, top_200_calm, single_stats,
                       custom_records, custom_summary, is_user_custom, filename):
    wb = openpyxl.Workbook()

    font_main_title = Font(name='Tahoma', size=11, bold=True, color='FFFFFF')
    font_sub_header = Font(name='Tahoma', size=9, bold=True, color='E2E8F0')
    font_data = Font(name='Tahoma', size=9, color='0F172A')
    font_bold = Font(name='Tahoma', size=9, bold=True, color='0F172A')
    font_profit = Font(name='Tahoma', size=9, bold=True, color='15803D')
    font_loss = Font(name='Tahoma', size=9, bold=True, color='B91C1C')
    font_stat_val = Font(name='Tahoma', size=11, bold=True, color='1E3A8A')
    font_stat_lbl = Font(name='Tahoma', size=9, color='475569')

    fill_period_header = PatternFill(start_color='1E293B', fill_type='solid')
    fill_sub_header = PatternFill(start_color='334155', fill_type='solid')
    fill_summary_header = PatternFill(start_color='047857', fill_type='solid')
    fill_summary_sub = PatternFill(start_color='065F46', fill_type='solid')
    fill_total_weight_hdr = PatternFill(start_color='4338CA', fill_type='solid')
    fill_zebra = PatternFill(start_color='F8FAFC', fill_type='solid')
    fill_card = PatternFill(start_color='F1F5F9', fill_type='solid')
    fill_highlight = PatternFill(start_color='DCFCE7', fill_type='solid')
    fill_total_weight = PatternFill(start_color='EEF2FF', fill_type='solid')

    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center')

    thin_border = Side(border_style='thin', color='CBD5E1')
    thick_border_left = Side(border_style='medium', color='64748B')
    thick_border_right = Side(border_style='medium', color='64748B')
    border_cell = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
    border_col_end = Border(left=thick_border_left, right=thin_border, top=thin_border, bottom=thin_border)
    border_total_weight = Border(left=thick_border_left, right=thick_border_right, top=thin_border, bottom=thin_border)

    N = len(MACRO_COMMODITIES)

    # -------------------------------------------------------------------------
    # شیت ۱: چرخش فصلی پویا دلاری (۱۱ دارایی)
    # -------------------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "چرخش فصلی پویا (۱۱ دارایی)"
    ws1.views.sheetView[0].rightToLeft = True

    ws1.merge_cells('A1:L1')
    ws1['A1'] = "گزارش چرخش فصلی پویا ماکرو دلاری — ۱۱ دارایی بین‌المللی (۴۵ دوره ۳ ماهه)"
    ws1['A1'].font = font_main_title; ws1['A1'].fill = fill_summary_header; ws1['A1'].alignment = align_center

    ws1.merge_cells('A2:L2')
    ws1['A2'] = "استراتژی: تخصیص فصلی ۱۰۰٪ به دارایی پیشتاز هر دوره (ستون تراز وزنی با فرمول صریح = ۱۰۰.۰٪)"
    ws1['A2'].font = font_bold; ws1['A2'].fill = fill_highlight; ws1['A2'].alignment = align_center

    stats_cards_dyn = [
        ("سود مرکب کل دلاری", f"{dyn_summary['cum_return']*100:+.2f}%"),
        ("میانگین سود فصلی", f"{dyn_summary['mean_return']*100:+.2f}%"),
        ("بیشترین افت (Max DD)", f"{dyn_summary['max_drawdown']*100:.2f}%"),
        ("شاخص شارپ دلاری", f"{dyn_summary['sharpe']:.2f}"),
        ("بهترین فصل دلاری", f"{dyn_summary['best_period']*100:+.2f}%"),
        ("بدترین فصل دلاری", f"{dyn_summary['worst_period']*100:+.2f}%"),
    ]
    _render_kpi_cards(ws1, stats_cards_dyn, font_stat_lbl, font_stat_val, fill_card, align_center)

    _build_macro_quarterly_sheet(
        ws1, periods_data, R, P_start, P_end, dyn_records, is_dynamic=True,
        font_main_title=font_main_title, font_sub_header=font_sub_header, font_data=font_data,
        font_bold=font_bold, font_profit=font_profit, font_loss=font_loss,
        fill_period_header=fill_period_header, fill_sub_header=fill_sub_header,
        fill_summary_header=fill_summary_header, fill_summary_sub=fill_summary_sub,
        fill_total_weight_hdr=fill_total_weight_hdr, fill_zebra=fill_zebra,
        fill_total_weight=fill_total_weight, align_center=align_center, align_right=align_right,
        border_cell=border_cell, border_col_end=border_col_end, border_total_weight=border_total_weight
    )

    # -------------------------------------------------------------------------
    # شیت ۲: پورتفوی با بیشترین بازدهی کل (۴۵ فصل)
    # -------------------------------------------------------------------------
    ws2 = wb.create_sheet(title="پورتفوی بیشترین بازدهی")
    ws2.views.sheetView[0].rightToLeft = True

    ws2.merge_cells('A1:L1')
    ws2['A1'] = "پورتفوی با وزن ثابت بیشترین بازدهی کل دلاری — ۴۵ دوره ۳ ماهه"
    ws2['A1'].font = font_main_title; ws2['A1'].fill = fill_summary_header; ws2['A1'].alignment = align_center

    w_desc_ret = " | ".join(f"{MACRO_COMMODITIES[i][1]}: {int(round(best_ret['weights'][i]*100))}%" 
                            for i in range(N) if best_ret['weights'][i] > 0)
    ws2.merge_cells('A2:L2')
    ws2['A2'] = f"ترکیب پورتفو: {w_desc_ret} (جمع اوزان در تمامی ردیف‌ها = ۱۰۰.۰٪)"
    ws2['A2'].font = font_bold; ws2['A2'].fill = fill_highlight; ws2['A2'].alignment = align_center

    stats_cards_ret = [
        ("سود مرکب کل دلاری", f"{best_ret['cum_return']*100:+.2f}%"),
        ("بازده سالانه (CAGR)", f"{best_ret['cagr']*100:+.2f}%"),
        ("بیشترین افت (Max DD)", f"{best_ret['max_drawdown']*100:.2f}%"),
        ("بیشترین ریزش یک فصل", f"{best_ret['worst_quarter']*100:.2f}%"),
        ("شاخص شارپ دلاری", f"{best_ret['sharpe']:.2f}"),
        ("شاخص کارایی کالمر", f"{best_ret['calmar']:.1f}"),
    ]
    _render_kpi_cards(ws2, stats_cards_ret, font_stat_lbl, font_stat_val, fill_card, align_center)

    _build_macro_quarterly_sheet(
        ws2, periods_data, R, P_start, P_end, fixed_ret_records, is_dynamic=False,
        font_main_title=font_main_title, font_sub_header=font_sub_header, font_data=font_data,
        font_bold=font_bold, font_profit=font_profit, font_loss=font_loss,
        fill_period_header=fill_period_header, fill_sub_header=fill_sub_header,
        fill_summary_header=fill_summary_header, fill_summary_sub=fill_summary_sub,
        fill_total_weight_hdr=fill_total_weight_hdr, fill_zebra=fill_zebra,
        fill_total_weight=fill_total_weight, align_center=align_center, align_right=align_right,
        border_cell=border_cell, border_col_end=border_col_end, border_total_weight=border_total_weight
    )

    # -------------------------------------------------------------------------
    # شیت ۳: پورتفوی بهینه بازدهی بالا با کمترین افت (Low Drawdown)
    # -------------------------------------------------------------------------
    ws3 = wb.create_sheet(title="پورتفوی بهینه بازدهی-افت")
    ws3.views.sheetView[0].rightToLeft = True

    ws3.merge_cells('A1:L1')
    ws3['A1'] = "پورتفوی بهینه با بیشترین بازدهی و کمترین Max Drawdown (سبد پایدار برتر شاخص کالمر)"
    ws3['A1'].font = font_main_title; ws3['A1'].fill = fill_summary_header; ws3['A1'].alignment = align_center

    w_desc_bal = " | ".join(f"{MACRO_COMMODITIES[i][1]}: {int(round(best_bal['weights'][i]*100))}%" 
                            for i in range(N) if best_bal['weights'][i] > 0)
    ws3.merge_cells('A2:L2')
    ws3['A2'] = f"ترکیب بهینه پایدار: {w_desc_bal} (تراز وزنی در هر ردیف = ۱۰۰.۰٪)"
    ws3['A2'].font = font_bold; ws3['A2'].fill = fill_highlight; ws3['A2'].alignment = align_center

    stats_cards_bal = [
        ("سود مرکب کل دلاری", f"{best_bal['cum_return']*100:+.2f}%"),
        ("بازده سالانه (CAGR)", f"{best_bal['cagr']*100:+.2f}%"),
        ("بیشترین افت (Max DD)", f"{best_bal['max_drawdown']*100:.2f}%"),
        ("بیشترین ریزش یک فصل", f"{best_bal['worst_quarter']*100:.2f}%"),
        ("شاخص شارپ دلاری", f"{best_bal['sharpe']:.2f}"),
        ("شاخص کارایی کالمر", f"{best_bal['calmar']:.2f}"),
    ]
    _render_kpi_cards(ws3, stats_cards_bal, font_stat_lbl, font_stat_val, fill_card, align_center)

    _build_macro_quarterly_sheet(
        ws3, periods_data, R, P_start, P_end, fixed_bal_records, is_dynamic=False,
        font_main_title=font_main_title, font_sub_header=font_sub_header, font_data=font_data,
        font_bold=font_bold, font_profit=font_profit, font_loss=font_loss,
        fill_period_header=fill_period_header, fill_sub_header=fill_sub_header,
        fill_summary_header=fill_summary_header, fill_summary_sub=fill_summary_sub,
        fill_total_weight_hdr=fill_total_weight_hdr, fill_zebra=fill_zebra,
        fill_total_weight=fill_total_weight, align_center=align_center, align_right=align_right,
        border_cell=border_cell, border_col_end=border_col_end, border_total_weight=border_total_weight
    )

    # -------------------------------------------------------------------------
    # شیت ۴: رده‌بندی ۲۰۰ پورتفوی برتر بر اساس بیشترین بازدهی (رنکینگ اصلی)
    # -------------------------------------------------------------------------
    ws4 = wb.create_sheet(title="رده‌بندی بر اساس بازدهی")
    ws4.views.sheetView[0].rightToLeft = True

    _build_ranking_sheet(
        ws4, top_200_ret, sort_name="بیشترین بازدهی دلاری کل (Cumulative Return)", is_calmar=False,
        font_main_title=font_main_title, font_sub_header=font_sub_header, font_data=font_data,
        font_bold=font_bold, font_profit=font_profit, font_loss=font_loss,
        fill_period_header=fill_period_header, fill_sub_header=fill_sub_header,
        fill_total_weight_hdr=fill_total_weight_hdr, fill_zebra=fill_zebra,
        fill_total_weight=fill_total_weight, align_center=align_center,
        border_cell=border_cell, border_total_weight=border_total_weight
    )

    # -------------------------------------------------------------------------
    # شیت ۵: رده‌بندی ۲۰۰ پورتفوی برتر بر اساس بیشترین بازدهی و کمترین Max Drawdown
    # -------------------------------------------------------------------------
    ws5 = wb.create_sheet(title="رده‌بندی بازدهی و کمترین افت")
    ws5.views.sheetView[0].rightToLeft = True

    _build_ranking_sheet(
        ws5, top_200_calm, sort_name="بهینه‌سازی ترکیبی بیشترین بازدهی و کمترین افت (Calmar Ratio)", is_calmar=True,
        font_main_title=font_main_title, font_sub_header=font_sub_header, font_data=font_data,
        font_bold=font_bold, font_profit=font_profit, font_loss=font_loss,
        fill_period_header=fill_period_header, fill_sub_header=fill_sub_header,
        fill_total_weight_hdr=fill_total_weight_hdr, fill_zebra=fill_zebra,
        fill_total_weight=fill_total_weight, align_center=align_center,
        border_cell=border_cell, border_total_weight=border_total_weight
    )

    # -------------------------------------------------------------------------
    # شیت ۶: عملکرد تک‌تک نمادهای ۱۱گانه (۱۱ سال کامل به دلار)
    # -------------------------------------------------------------------------
    ws6 = wb.create_sheet(title="عملکرد نمادهای منفرد (۱۱ سال)")
    ws6.views.sheetView[0].rightToLeft = True

    ws6.merge_cells('A1:K1')
    ws6['A1'] = "عملکرد تک‌تک دارایی‌های ۱۱گانه در بازار جهانی به دلار (۲۰۱۵ تا ۲۰۲۶ - ۴۵ دوره ۳ ماهه)"
    ws6['A1'].font = font_main_title; ws6['A1'].fill = fill_period_header; ws6['A1'].alignment = align_center

    headers_single = [
        "نام دارایی", "تیکر جهانی", "قیمت اول ۲۰۱۵ ($)", "قیمت پایان ۲۰۲۶ ($)",
        "بازدهی کل دلاری (٪)", "میانگین فصلی (٪)", "انحراف معیار (٪)",
        "Max Drawdown (عمیق‌ترین افت ٪)", "بیشترین ریزش یک فصل (٪)", "بهترین فصل (٪)", "شاخص شارپ"
    ]
    for col_idx, h_text in enumerate(headers_single, 1):
        cell = ws6.cell(2, col_idx, h_text)
        cell.font = font_sub_header; cell.fill = fill_sub_header; cell.alignment = align_center; cell.border = border_cell

    for idx, stat in enumerate(single_stats, 1):
        r_num = idx + 2
        is_even = (idx % 2 == 0)
        r_fill = fill_zebra if is_even else None

        vals = [
            (stat['name'], font_bold, align_center, None),
            (stat['ticker'], font_data, align_center, None),
            (stat['p_start'], font_data, align_right, '$#,##0.00'),
            (stat['p_end'], font_data, align_right, '$#,##0.00'),
            (stat['cum_return'], font_profit if stat['cum_return'] > 0 else font_loss, align_center, '+0.00%;-0.00%;0.00%'),
            (stat['mean_q'], font_profit if stat['mean_q'] > 0 else font_loss, align_center, '+0.00%;-0.00%;0.00%'),
            (stat['std_q'], font_data, align_center, '0.00%'),
            (stat['max_drawdown'], font_loss, align_center, '+0.00%;-0.00%;0.00%'),
            (stat['worst_q'], font_loss, align_center, '+0.00%;-0.00%;0.00%'),
            (stat['best_q'], font_profit, align_center, '+0.00%;-0.00%;0.00%'),
            (stat['sharpe'], font_bold, align_center, '0.00'),
        ]

        for col_idx, (val, f_style, a_style, num_fmt) in enumerate(vals, 1):
            cell = ws6.cell(r_num, col_idx, val)
            cell.font = f_style; cell.alignment = a_style; cell.border = border_cell
            if num_fmt: cell.number_format = num_fmt
            if r_fill: cell.fill = r_fill

    for col in ws6.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws6.column_dimensions[col_letter].width = max(max_len + 4, 13)

    # -------------------------------------------------------------------------
    # شیت ۷: پورتفوی سفارشی شما (Custom Portfolio Backtest)
    # -------------------------------------------------------------------------
    ws7 = wb.create_sheet(title="پورتفوی سفارشی شما")
    ws7.views.sheetView[0].rightToLeft = True

    ws7.merge_cells('A1:L1')
    if is_user_custom:
        ws7['A1'] = "گزارش بک‌تست تاریخی پورتفوی سفارشی پیشنهادی شما — ۴۵ دوره ۳ ماهه (۱۱ سال کامل)"
    else:
        ws7['A1'] = "گزارش بک‌تست پورتفوی نمونه (برای تست ترکیب دلخواه: اجرای اسکریپت با آرگومان --custom)"
    ws7['A1'].font = font_main_title; ws7['A1'].fill = PatternFill(start_color='4F46E5', fill_type='solid'); ws7['A1'].alignment = align_center

    w_desc_custom = " | ".join(f"{MACRO_COMMODITIES[i][1]}: {int(round(custom_summary['weights'][i]*100))}%" 
                               for i in range(N) if custom_summary['weights'][i] > 0)
    ws7.merge_cells('A2:L2')
    ws7['A2'] = f"ترکیب پورتفو: {w_desc_custom} (جمع اوزان در تمامی ردیف‌ها = ۱۰۰.۰٪)"
    ws7['A2'].font = font_bold; ws7['A2'].fill = fill_highlight; ws7['A2'].alignment = align_center

    stats_cards_custom = [
        ("سود مرکب کل دلاری", f"{custom_summary['cum_return']*100:+.2f}%"),
        ("بازده سالانه (CAGR)", f"{custom_summary['cagr']*100:+.2f}%"),
        ("بیشترین افت (Max DD)", f"{custom_summary['max_drawdown']*100:.2f}%"),
        ("بیشترین ریزش یک فصل", f"{custom_summary['worst_quarter']*100:.2f}%"),
        ("شاخص شارپ دلاری", f"{custom_summary['sharpe']:.2f}"),
        ("شاخص کارایی کالمر", f"{custom_summary['calmar']:.2f}"),
    ]
    _render_kpi_cards(ws7, stats_cards_custom, font_stat_lbl, font_stat_val, fill_card, align_center)

    _build_macro_quarterly_sheet(
        ws7, periods_data, R, P_start, P_end, custom_records, is_dynamic=False,
        font_main_title=font_main_title, font_sub_header=font_sub_header, font_data=font_data,
        font_bold=font_bold, font_profit=font_profit, font_loss=font_loss,
        fill_period_header=fill_period_header, fill_sub_header=fill_sub_header,
        fill_summary_header=fill_summary_header, fill_summary_sub=fill_summary_sub,
        fill_total_weight_hdr=fill_total_weight_hdr, fill_zebra=fill_zebra,
        fill_total_weight=fill_total_weight, align_center=align_center, align_right=align_right,
        border_cell=border_cell, border_col_end=border_col_end, border_total_weight=border_total_weight
    )

    # ذخیره نهایی با مدیریت هوشمند فایل قفل‌شده در اکسل
    try:
        wb.save(filename)
    except PermissionError:
        alt_filename = filename.replace('.xlsx', '_v2.xlsx')
        wb.save(alt_filename)
        print(f"\n⚠️ فایل '{filename}' توسط نرم‌افزار اکسل قفل شده است.")
        print(f"✅ گزارش جدید با نام جایگزین ذخیره شد: {os.path.abspath(alt_filename)}")


def _render_kpi_cards(ws, cards, font_lbl, font_val, fill_card, align_center):
    for col_i, (lbl, val) in enumerate(cards):
        c1 = col_i * 2 + 1
        ws.merge_cells(start_row=3, start_column=c1, end_row=3, end_column=c1 + 1)
        ws.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c1 + 1)
        cl = ws.cell(3, c1, lbl); cl.font = font_lbl; cl.fill = fill_card; cl.alignment = align_center
        cv = ws.cell(4, c1, val); cv.font = font_val; cv.fill = fill_card; cv.alignment = align_center


def _build_macro_quarterly_sheet(ws, periods_data, R, P_start, P_end, records, is_dynamic,
                                 font_main_title, font_sub_header, font_data, font_bold,
                                 font_profit, font_loss, fill_period_header, fill_sub_header,
                                 fill_summary_header, fill_summary_sub, fill_total_weight_hdr,
                                 fill_zebra, fill_total_weight, align_center, align_right,
                                 border_cell, border_col_end, border_total_weight):
    T = len(periods_data)
    N = len(MACRO_COMMODITIES)

    ws.merge_cells('A6:D6')
    ws['A6'] = "مشخصات دوره زمانی"
    ws['A6'].font = font_main_title; ws['A6'].fill = fill_period_header; ws['A6'].alignment = align_center

    meta_cols = ["دوره", "تاریخ میلادی", "تاریخ شمسی", "طول دوره (روز)"]
    for idx, mc in enumerate(meta_cols, 1):
        cell = ws.cell(7, idx, mc)
        cell.font = font_sub_header; cell.fill = fill_sub_header; cell.alignment = align_center; cell.border = border_cell

    col_cur = 5
    comm_col_map = {}
    weight_cols = []
    for cid, cname, ticker, color_hex in MACRO_COMMODITIES:
        s_c = col_cur
        e_c = col_cur + 3
        comm_col_map[cid] = s_c
        weight_cols.append(s_c)

        ws.merge_cells(start_row=6, start_column=s_c, end_row=6, end_column=e_c)
        h_c = ws.cell(6, s_c, f"{cname} [{ticker}]")
        h_c.font = font_main_title
        h_c.fill = PatternFill(start_color=color_hex, fill_type='solid')
        h_c.alignment = align_center

        sub_labels = ["وزن پورتفو", "تغییر قیمت دلاری", "قیمت اول ($)", "قیمت پایان ($)"]
        for o_i, sl in enumerate(sub_labels):
            cell = ws.cell(7, s_c + o_i, sl)
            cell.font = font_sub_header; cell.fill = fill_sub_header; cell.alignment = align_center
            cell.border = border_cell if o_i < 3 else border_col_end
        col_cur += 4

    # ستون صریح جمع وزن‌ها با فرمول = ۱۰۰.۰٪
    tot_w_col = col_cur
    ws.cell(6, tot_w_col, "تراز وزنی")
    ws.cell(6, tot_w_col).font = font_main_title; ws.cell(6, tot_w_col).fill = fill_total_weight_hdr; ws.cell(6, tot_w_col).alignment = align_center

    cell_tot_hdr = ws.cell(7, tot_w_col, "جمع وزن‌ها (۱۰۰٪)")
    cell_tot_hdr.font = font_sub_header; cell_tot_hdr.fill = fill_total_weight_hdr; cell_tot_hdr.alignment = align_center; cell_tot_hdr.border = border_total_weight
    col_cur += 1

    sum_col = col_cur
    extra_cols = 4 if is_dynamic else 3
    ws.merge_cells(start_row=6, start_column=sum_col, end_row=6, end_column=sum_col + extra_cols - 1)
    sh_c = ws.cell(6, sum_col, "عملکرد پورتفوی دلاری")
    sh_c.font = font_main_title; sh_c.fill = fill_summary_header; sh_c.alignment = align_center

    if is_dynamic:
        perf_sub = ["نماد منتخب فصل", "بازدهی فصل پورتفو (٪)", "ارزش سرمایه دلاری ($)", "رشد کل دلاری (٪)"]
    else:
        perf_sub = ["بازدهی فصل پورتفو (٪)", "ارزش سرمایه دلاری ($)", "رشد کل دلاری (٪)"]

    for idx, p_label in enumerate(perf_sub):
        sc = ws.cell(7, sum_col + idx, p_label)
        sc.font = font_sub_header; sc.fill = fill_summary_sub; sc.alignment = align_center; sc.border = border_cell

    cur_r = 8
    for t in range(T):
        p_info = periods_data[t]
        rec = records[t]
        is_even = (t % 2 == 1)
        r_fill = fill_zebra if is_even else None

        c_p = ws.cell(cur_r, 1, f"دوره {p_info['period_num']:02d}")
        c_gr = ws.cell(cur_r, 2, f"{p_info['greg_start']} -> {p_info['greg_end']}")
        c_sh = ws.cell(cur_r, 3, f"{p_info['jalali_start']} -> {p_info['jalali_end']}")
        c_dy = ws.cell(cur_r, 4, p_info['days'])

        for c in (c_p, c_gr, c_sh, c_dy):
            c.font = font_bold if c == c_p else font_data; c.alignment = align_center; c.border = border_cell
            if r_fill: c.fill = r_fill

        for c_idx, (cid, cname, ticker, _) in enumerate(MACRO_COMMODITIES):
            st_c = comm_col_map[cid]
            w_val = rec['weights'][c_idx]
            r_val = R[t, c_idx]
            p1_val = P_start[t, c_idx]
            p2_val = P_end[t, c_idx]

            cw = ws.cell(cur_r, st_c, w_val)
            cw.number_format = '0.0%'
            cw.font = font_bold if w_val > 0 else font_data
            cw.alignment = align_center
            if w_val > 0:
                cw.fill = PatternFill(start_color='DCFCE7', fill_type='solid')

            cr = ws.cell(cur_r, st_c + 1, r_val)
            cr.number_format = '+0.00%;-0.00%;0.00%'
            cr.font = font_profit if r_val > 0 else (font_loss if r_val < 0 else font_data)
            cr.alignment = align_center

            cp1 = ws.cell(cur_r, st_c + 2, p1_val)
            cp1.number_format = '$#,##0.00' if p1_val >= 1 else '$#,##0.0000'
            cp1.font = font_data; cp1.alignment = align_right

            cp2 = ws.cell(cur_r, st_c + 3, p2_val)
            cp2.number_format = '$#,##0.00' if p2_val >= 1 else '$#,##0.0000'
            cp2.font = font_data; cp2.alignment = align_right

            for o_i, c in enumerate((cw, cr, cp1, cp2)):
                c.border = border_cell if o_i < 3 else border_col_end
                if r_fill and c.fill.fill_type is None:
                    c.fill = r_fill

        # فرمول جمع وزن‌های هر ردیف = ۱۰۰.۰٪
        sum_args = ",".join(f"{get_column_letter(w_col)}{cur_r}" for w_col in weight_cols)
        c_tot = ws.cell(cur_r, tot_w_col, f"=SUM({sum_args})")
        c_tot.number_format = '0.0%'
        c_tot.font = font_bold
        c_tot.alignment = align_center
        c_tot.fill = fill_total_weight
        c_tot.border = border_total_weight

        p_ret = rec['period_ret']
        cum_cap = rec['cum_capital']
        cum_ret = rec['cum_return']

        col_offset = 0
        if is_dynamic:
            c_chosen = ws.cell(cur_r, sum_col + col_offset, rec['chosen_name'])
            c_chosen.font = font_bold; c_chosen.alignment = align_center; c_chosen.border = border_cell
            if r_fill: c_chosen.fill = r_fill
            col_offset += 1

        c_pret = ws.cell(cur_r, sum_col + col_offset, p_ret)
        c_pret.number_format = '+0.00%;-0.00%;0.00%'
        c_pret.font = font_profit if p_ret > 0 else (font_loss if p_ret < 0 else font_data)
        c_pret.alignment = align_center; c_pret.border = border_cell
        if r_fill: c_pret.fill = r_fill
        col_offset += 1

        c_cap = ws.cell(cur_r, sum_col + col_offset, cum_cap)
        c_cap.number_format = '$#,##0.00'
        c_cap.font = font_bold; c_cap.alignment = align_right; c_cap.border = border_cell
        if r_fill: c_cap.fill = r_fill
        col_offset += 1

        c_cum = ws.cell(cur_r, sum_col + col_offset, cum_ret)
        c_cum.number_format = '+0.00%;-0.00%;0.00%'
        c_cum.font = font_bold; c_cum.alignment = align_center; c_cum.border = border_cell
        if r_fill: c_cum.fill = r_fill

        cur_r += 1

    ws.column_dimensions['A'].width = 11
    ws.column_dimensions['B'].width = 24
    ws.column_dimensions['C'].width = 24
    ws.column_dimensions['D'].width = 15
    for w_col in weight_cols:
        ws.column_dimensions[get_column_letter(w_col)].width = 12
        ws.column_dimensions[get_column_letter(w_col + 1)].width = 15
        ws.column_dimensions[get_column_letter(w_col + 2)].width = 14
        ws.column_dimensions[get_column_letter(w_col + 3)].width = 14
    ws.column_dimensions[get_column_letter(tot_w_col)].width = 17
    ws.column_dimensions[get_column_letter(sum_col)].width = 24 if is_dynamic else 18
    ws.column_dimensions[get_column_letter(sum_col + 1)].width = 18
    ws.column_dimensions[get_column_letter(sum_col + 2)].width = 18
    if is_dynamic:
        ws.column_dimensions[get_column_letter(sum_col + 3)].width = 18


def _build_ranking_sheet(ws, top_list, sort_name, is_calmar, font_main_title, font_sub_header, font_data,
                         font_bold, font_profit, font_loss, fill_period_header, fill_sub_header,
                         fill_total_weight_hdr, fill_zebra, fill_total_weight, align_center,
                         border_cell, border_total_weight):
    N = len(MACRO_COMMODITIES)

    max_c_letter = 'S' if is_calmar else 'R'
    ws.merge_cells(f'A1:{max_c_letter}1')
    ws['A1'] = f"رده‌بندی ۲۰۰ پورتفوی برتر ماکرو دلاری — معیار رتبه‌بندی: {sort_name}"
    ws['A1'].font = font_main_title; ws['A1'].fill = fill_period_header; ws['A1'].alignment = align_center

    if is_calmar:
        headers = [
            "رتبه",
            "سکه طلا", "طلا ۲۴", "نقره", "مس", "روی", "سنگ آهن", "قیر/نفت", "بیت‌کوین", "اتریوم", "S&P 500", "تتر",
            "جمع وزن‌ها (۱۰۰٪)",
            "شاخص کالمر (بازده سالانه/افت)", "Max Drawdown (بیشترین افت کل)", "بازده مرکب سالانه (CAGR ٪)",
            "سود کل دلاری (٪)", "بیشترین ریزش یک فصل", "شاخص شارپ"
        ]
    else:
        headers = [
            "رتبه",
            "سکه طلا", "طلا ۲۴", "نقره", "مس", "روی", "سنگ آهن", "قیر/نفت", "بیت‌کوین", "اتریوم", "S&P 500", "تتر",
            "جمع وزن‌ها (۱۰۰٪)",
            "سود کل دلاری (٪)", "بازده مرکب سالانه (CAGR ٪)", "Max Drawdown (بیشترین افت کل)",
            "بیشترین ریزش یک فصل", "شاخص شارپ"
        ]

    for col_idx, h_text in enumerate(headers, 1):
        cell = ws.cell(2, col_idx, h_text)
        cell.font = font_sub_header
        cell.fill = fill_total_weight_hdr if "جمع" in h_text else fill_sub_header
        cell.alignment = align_center; cell.border = border_cell

    for rank_i, p_info in enumerate(top_list, 1):
        r_num = rank_i + 2
        is_even = (rank_i % 2 == 0)
        r_fill = fill_zebra if is_even else None

        c_rk = ws.cell(r_num, 1, f"#{rank_i:02d}")
        c_rk.font = font_bold; c_rk.alignment = align_center; c_rk.border = border_cell
        if r_fill: c_rk.fill = r_fill

        for w_i in range(N):
            w_val = p_info['weights'][w_i]
            cw = ws.cell(r_num, w_i + 2, w_val)
            cw.number_format = '0%'
            cw.font = font_bold if w_val > 0 else font_data
            cw.alignment = align_center; cw.border = border_cell
            if w_val > 0:
                cw.fill = PatternFill(start_color='DCFCE7', fill_type='solid')
            elif r_fill:
                cw.fill = r_fill

        # ستون جمع وزن‌ها = SUM(ColB:ColL)
        c_tot = ws.cell(r_num, 13, f"=SUM(B{r_num}:L{r_num})")
        c_tot.number_format = '0.0%'
        c_tot.font = font_bold; c_tot.alignment = align_center
        c_tot.fill = fill_total_weight; c_tot.border = border_total_weight

        if is_calmar:
            # ۱. شاخص کالمر
            c_calm = ws.cell(r_num, 14, p_info['calmar'])
            c_calm.number_format = '0.00'
            c_calm.font = font_bold; c_calm.alignment = align_center; c_calm.border = border_cell
            if r_fill: c_calm.fill = r_fill

            # ۲. Max Drawdown
            c_mdd = ws.cell(r_num, 15, p_info['max_drawdown'])
            c_mdd.number_format = '+0.00%;-0.00%;0.00%'
            c_mdd.font = font_loss; c_mdd.alignment = align_center; c_mdd.border = border_cell
            if r_fill: c_mdd.fill = r_fill

            # ۳. CAGR
            c_cagr = ws.cell(r_num, 16, p_info['cagr'])
            c_cagr.number_format = '+0.00%;-0.00%;0.00%'
            c_cagr.font = font_profit; c_cagr.alignment = align_center; c_cagr.border = border_cell
            if r_fill: c_cagr.fill = r_fill

            # ۴. سود کل دلاری
            c_cum = ws.cell(r_num, 17, p_info['cum_return'])
            c_cum.number_format = '+0.00%;-0.00%;0.00%'
            c_cum.font = font_profit; c_cum.alignment = align_center; c_cum.border = border_cell
            if r_fill: c_cum.fill = r_fill

            # ۵. بیشترین ریزش یک فصل
            c_wq = ws.cell(r_num, 18, p_info['worst_quarter'])
            c_wq.number_format = '+0.00%;-0.00%;0.00%'
            c_wq.font = font_loss; c_wq.alignment = align_center; c_wq.border = border_cell
            if r_fill: c_wq.fill = r_fill

            # ۶. شارپ
            c_sh = ws.cell(r_num, 19, p_info['sharpe'])
            c_sh.number_format = '0.00'
            c_sh.font = font_bold; c_sh.alignment = align_center; c_sh.border = border_cell
            if r_fill: c_sh.fill = r_fill
        else:
            # ۱. سود کل دلاری
            c_cum = ws.cell(r_num, 14, p_info['cum_return'])
            c_cum.number_format = '+0.00%;-0.00%;0.00%'
            c_cum.font = font_profit; c_cum.alignment = align_center; c_cum.border = border_cell
            if r_fill: c_cum.fill = r_fill

            # ۲. CAGR
            c_cagr = ws.cell(r_num, 15, p_info['cagr'])
            c_cagr.number_format = '+0.00%;-0.00%;0.00%'
            c_cagr.font = font_profit; c_cagr.alignment = align_center; c_cagr.border = border_cell
            if r_fill: c_cagr.fill = r_fill

            # ۳. Max Drawdown
            c_mdd = ws.cell(r_num, 16, p_info['max_drawdown'])
            c_mdd.number_format = '+0.00%;-0.00%;0.00%'
            c_mdd.font = font_loss; c_mdd.alignment = align_center; c_mdd.border = border_cell
            if r_fill: c_mdd.fill = r_fill

            # ۴. بیشترین ریزش یک فصل
            c_wq = ws.cell(r_num, 17, p_info['worst_quarter'])
            c_wq.number_format = '+0.00%;-0.00%;0.00%'
            c_wq.font = font_loss; c_wq.alignment = align_center; c_wq.border = border_cell
            if r_fill: c_wq.fill = r_fill

            # ۵. شارپ
            c_sh = ws.cell(r_num, 18, p_info['sharpe'])
            c_sh.number_format = '0.00'
            c_sh.font = font_bold; c_sh.alignment = align_center; c_sh.border = border_cell
            if r_fill: c_sh.fill = r_fill

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
    ws.column_dimensions['M'].width = 16
    ws.column_dimensions['N'].width = 22
    ws.column_dimensions['O'].width = 22
    ws.column_dimensions['P'].width = 22
    if is_calmar:
        ws.column_dimensions['Q'].width = 22
        ws.column_dimensions['R'].width = 18


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="BKV Global Macro Multi-Asset Portfolio Optimizer")
    parser.add_argument('--custom', '-c', action='store_true', help="درخواست ورودی تعاملی درصدهای سفارشی در ترمینال برای بک‌تست")
    parser.add_argument('--weights', '-w', type=str, default=None, help="وزن‌های سفارشی به صورت کاما یا key=val (مثال: 'gold_995=30,btc=10,eth=15,spx=20,usdt=25')")
    parser.add_argument('--no-interactive', action='store_true', help="اجرای کاملاً خودکار بدون پرامپت ترمینال")
    cli_args = parser.parse_args()

    run_macro_global_optimization(cli_args)
